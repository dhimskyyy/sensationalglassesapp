from __future__ import annotations

import re
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.shared import Pt


SRC_DOCX = Path("/Users/mhmmddhimas/SKRIPSIIII/ta2/FIX_COPYRevisi_TA2_Mohammad Dhimas Afrizal_2211104023.docx")
SRC_XLSX = Path("/Users/mhmmddhimas/Downloads/User Acceptance Testing (UAT) Sensational Glasses App (Jawaban).xlsx")
OUT_DIR = Path("/Users/mhmmddhimas/SKRIPSIIII/sensationalglassesapp/outputs/skripsi_uat_real")
OUT_DOCX = OUT_DIR / "FIX_COPYRevisi_TA2_Mohammad Dhimas Afrizal_2211104023_UAT_REAL.docx"


LIKERT_CODES = [
    "A1", "A2", "A3",
    "B1", "B2", "B3",
    "C1", "C2", "C3",
    "D1", "D2",
    "E1", "E2", "E3", "E4",
]

VARIABLES = [
    ("Kemudahan Penggunaan Aplikasi", ["A1", "A2", "A3"]),
    ("Kejelasan Informasi dan Monitoring Perangkat", ["B1", "B2", "B3"]),
    ("Peta, Navigasi, dan Notifikasi", ["C1", "C2", "C3"]),
    ("Kontrol Alarm", ["D1", "D2"]),
    ("Kesesuaian dan Kelayakan Aplikasi Pendamping", ["E1", "E2", "E3", "E4"]),
]


def clean_text(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    text = text.replace("aplokbermanfaat", "aplikasi bermanfaat")
    return text


def fmt_dec(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


def fmt_pct(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return f"{int(round(value))}%"
    return f"{value:.2f}".replace(".", ",") + "%"


def category(pct: float) -> str:
    if pct <= 20:
        return "Sangat Kurang Baik"
    if pct <= 40:
        return "Kurang Baik"
    if pct <= 60:
        return "Cukup Baik"
    if pct <= 80:
        return "Baik"
    return "Sangat Baik"


def load_uat():
    df = pd.read_excel(SRC_XLSX)
    df.columns = [clean_text(c) for c in df.columns]
    likert_cols = list(df.columns[6:21])
    open_cols = list(df.columns[21:25])

    rows = []
    total_actual = 0
    pct_by_code = {}
    question_by_code = {}
    for code, col in zip(LIKERT_CODES, likert_cols):
        vals = pd.to_numeric(df[col], errors="coerce")
        total = int(vals.sum())
        mean = float(vals.mean())
        pct = mean / 5 * 100
        total_actual += total
        pct_by_code[code] = pct
        question_by_code[code] = clean_text(col)
        rows.append([code, str(total), fmt_dec(mean), fmt_pct(pct), category(pct)])

    max_score = len(df) * len(likert_cols) * 5
    final_pct = total_actual / max_score * 100

    profile_rows = []
    for idx, row in df.iterrows():
        profile_rows.append([
            str(idx + 1),
            f"R{idx + 1}",
            clean_text(row["Instansi"]),
            clean_text(row["Jabatan/Peran"]),
            clean_text(row["Apakah pernah mengajar atau mendampingi siswa tunanetra?"]),
            clean_text(row["Berapa lama pengalaman anda mendampingi penyandang tunanetra"]),
        ])

    variable_rows = []
    no = 1
    for var, codes in VARIABLES:
        for code in codes:
            variable_rows.append([str(no), var if code == codes[0] else "", question_by_code[code], code])
            no += 1

    avg_rows = []
    for i, (var, codes) in enumerate(VARIABLES, 1):
        avg = sum(pct_by_code[c] for c in codes) / len(codes)
        avg_rows.append([str(i), var, ", ".join(codes), fmt_pct(avg), category(avg)])

    open_summary_rows = [
        [
            "1",
            "Fitur yang paling membantu",
            "Responden menilai fitur monitoring, pelacakan atau pemantauan alat, serta kontrol alarm sebagai fitur yang paling membantu dalam mendukung proses pendampingan.",
        ],
        [
            "2",
            "Bagian yang masih membingungkan atau perlu diperbaiki",
            "Sebagian besar responden menyatakan aplikasi sudah cukup baik dan tidak membingungkan. Namun, terdapat masukan bahwa bagian pengisian kredensial alat masih perlu diperjelas.",
        ],
        [
            "3",
            "Fitur tambahan yang diperlukan",
            "Mayoritas responden tidak mengusulkan fitur tambahan khusus dan menilai fitur yang tersedia sudah cukup untuk kebutuhan implementasi terbatas.",
        ],
        [
            "4",
            "Saran atau masukan lain",
            "Responden memberikan masukan positif bahwa aplikasi bermanfaat bagi pendamping penyandang tunanetra, sudah baik, dan dapat dilanjutkan menuju penyempurnaan akhir.",
        ],
    ]

    appendix_rows = []
    for i, row in df.iterrows():
        appendix_rows.append([
            f"R{i+1}",
            clean_text(row["Jabatan/Peran"]),
            clean_text(row["Berapa lama pengalaman anda mendampingi penyandang tunanetra"]),
            "; ".join(str(int(pd.to_numeric(row[col], errors="coerce"))) for col in likert_cols),
            clean_text(row[open_cols[0]]),
            clean_text(row[open_cols[1]]),
            clean_text(row[open_cols[2]]),
            clean_text(row[open_cols[3]]),
        ])

    return {
        "n": len(df),
        "profile_rows": profile_rows,
        "variable_rows": variable_rows,
        "rekap_rows": rows,
        "avg_rows": avg_rows,
        "open_summary_rows": open_summary_rows,
        "appendix_rows": appendix_rows,
        "total_actual": total_actual,
        "max_score": max_score,
        "final_pct": final_pct,
        "final_category": category(final_pct),
        "likert_cols": likert_cols,
    }


def set_para_text(paragraph, text: str):
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def find_para(doc: Document, text: str, start: int = 0, style: str | None = None):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if p.text.strip() == text:
            if style is None or p.style.name == style:
                return i, p
    raise ValueError(f"Paragraph not found: {text!r}")


def delete_between(start_para, end_para):
    current = start_para._p.getnext()
    while current is not None and current is not end_para._p:
        nxt = current.getnext()
        current.getparent().remove(current)
        current = nxt


def add_paragraph_before(doc: Document, before_el, text: str = "", style: str | None = None, align=None):
    p = doc.add_paragraph(text)
    if style:
        p.style = style
    if align is not None:
        p.alignment = align
    p._p.getparent().remove(p._p)
    before_el.addprevious(p._p)
    return p


def add_table_before(doc: Document, before_el, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)

    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=80, start=80, bottom=80, end=80)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)
    set_table_borders(table)

    table._tbl.getparent().remove(table._tbl)
    before_el.addprevious(table._tbl)
    return table


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def add_caption_before(doc: Document, before_el, text: str):
    p = add_paragraph_before(doc, before_el, text, style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.bold = True
    return p


def add_uat_section(doc: Document, before_el, data):
    n = data["n"]
    final_pct = fmt_dec(data["final_pct"])
    final_cat = data["final_category"]
    total = data["total_actual"]
    max_score = data["max_score"]

    paras = [
        "User Acceptance Testing (UAT) dilakukan untuk memperoleh penilaian dari sisi pendamping siswa tunanetra terhadap kesesuaian aplikasi Sensational Glasses App setelah aplikasi dicoba secara langsung. Pada tahap ini, responden mencoba aplikasi yang telah tersedia melalui Google Play Store, kemudian memberikan penilaian melalui kuesioner Google Form berbasis skala Likert 1-5 dan pertanyaan terbuka.",
        "Responden UAT pada penelitian ini adalah guru atau tenaga pendidik di SLB Negeri Purwokerto yang memiliki pengalaman mengajar atau mendampingi siswa tunanetra. Identitas pribadi responden tidak ditampilkan dalam laporan penelitian dan digantikan dengan kode responden R1 sampai R5. Profil responden UAT ditunjukkan pada Tabel 4.10.",
    ]
    for p in paras:
        add_paragraph_before(doc, before_el, "\t" + p)

    add_caption_before(doc, before_el, "Tabel 4.10 Profil Responden UAT")
    add_table_before(
        doc,
        before_el,
        ["No", "Kode Responden", "Instansi", "Jabatan/Peran", "Pernah Mengajar atau Mendampingi Siswa Tunanetra", "Lama Pengalaman"],
        data["profile_rows"],
        widths=[0.35, 0.75, 1.15, 0.9, 2.05, 1.0],
    )
    add_paragraph_before(doc, before_el, "\tBerdasarkan Tabel 4.10, seluruh responden merupakan guru di SLB Negeri Purwokerto dan memiliki pengalaman mengajar atau mendampingi siswa tunanetra. Dengan demikian, responden dinilai relevan untuk memberikan penilaian terhadap kesesuaian aplikasi Sensational Glasses App dari sisi pendamping siswa tunanetra.")
    add_paragraph_before(doc, before_el, "\tInstrumen UAT disusun berdasarkan beberapa variabel penilaian yang disesuaikan dengan fitur dan tujuan aplikasi. Daftar variabel dan pertanyaan kuesioner UAT ditunjukkan pada Tabel 4.11.")
    add_caption_before(doc, before_el, "Tabel 4.11 Daftar Variabel dan Pertanyaan Kuesioner UAT")
    add_table_before(doc, before_el, ["No", "Variabel", "Pertanyaan", "Kode"], data["variable_rows"], widths=[0.35, 1.55, 3.65, 0.45])
    add_paragraph_before(doc, before_el, "\tTabel 4.11 menunjukkan bahwa kuesioner UAT terdiri dari lima variabel utama, yaitu kemudahan penggunaan aplikasi, kejelasan informasi dan monitoring perangkat, peta, navigasi, dan notifikasi, kontrol alarm, serta kesesuaian dan kelayakan aplikasi pendamping. Kelima variabel tersebut digunakan untuk menilai apakah aplikasi telah sesuai dengan kebutuhan pendamping dalam proses monitoring dan tracking perangkat Sensational Glasses.")
    add_paragraph_before(doc, before_el, "\tPengolahan hasil UAT dilakukan dengan menghitung total skor, nilai rata-rata, dan persentase pada setiap pernyataan. Rumus yang digunakan adalah sebagai berikut:")
    add_paragraph_before(doc, before_el, "Mean = Total Skor Pernyataan / Jumlah Responden", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, "Persentase = (Mean / Bobot Maksimum) x 100%", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, "Keterangan:")
    for item in [
        "Total Skor Pernyataan adalah jumlah seluruh bobot jawaban responden pada satu pernyataan.",
        "Jumlah Responden adalah jumlah responden yang mengisi kuesioner UAT.",
        "Bobot Maksimum adalah nilai tertinggi pada skala Likert, yaitu 5.",
    ]:
        add_paragraph_before(doc, before_el, item, style="List Paragraph")
    add_paragraph_before(doc, before_el, "\tRekapitulasi hasil perhitungan UAT ditunjukkan pada Tabel 4.12.")
    add_caption_before(doc, before_el, "Tabel 4.12 Rekapitulasi Hasil User Acceptance Testing (UAT)")
    add_table_before(doc, before_el, ["Kode", "Total Skor", "Mean", "Persentase", "Keterangan"], data["rekap_rows"], widths=[0.8, 1.1, 0.9, 1.1, 1.35])
    add_paragraph_before(doc, before_el, "\tBerdasarkan Tabel 4.12, seluruh pernyataan UAT memperoleh persentase antara 72% sampai 80% dan termasuk dalam kategori Baik. Persentase tertinggi sebesar 80% terdapat pada pernyataan B3, C1, D1, E2, E3, dan E4. Hal ini menunjukkan bahwa fitur monitoring, fitur peta, kontrol alarm, kesesuaian informasi, potensi pendampingan, dan kelayakan aplikasi memperoleh penilaian baik dari responden.")
    add_paragraph_before(doc, before_el, "\tSelain perhitungan pada setiap pernyataan, hasil UAT juga dihitung berdasarkan rata-rata setiap variabel penilaian. Rekapitulasi rata-rata per variabel ditunjukkan pada Tabel 4.13.")
    add_caption_before(doc, before_el, "Tabel 4.13 Rata-Rata Hasil UAT Per Variabel")
    add_table_before(doc, before_el, ["No", "Variabel", "Kode Pernyataan", "Rata-Rata Persentase", "Keterangan"], data["avg_rows"], widths=[0.35, 2.3, 1.1, 1.3, 1.0])
    add_paragraph_before(doc, before_el, "\tTabel 4.13 menunjukkan bahwa seluruh variabel penilaian memperoleh kategori Baik. Variabel dengan nilai tertinggi adalah kesesuaian dan kelayakan aplikasi pendamping dengan persentase 79%, sedangkan variabel kemudahan penggunaan aplikasi memperoleh nilai 74,67%. Hasil ini menunjukkan bahwa aplikasi dinilai sesuai untuk mendukung pendamping dalam melakukan pemantauan, pelacakan, penerimaan notifikasi, penggunaan peta dan navigasi, serta kontrol alarm pada perangkat Sensational Glasses. Namun, aspek kemudahan penggunaan tetap menjadi perhatian karena memiliki nilai paling rendah dibandingkan variabel lainnya.")
    add_paragraph_before(doc, before_el, "\tNilai akhir UAT digunakan untuk mengetahui tingkat penerimaan aplikasi secara keseluruhan berdasarkan seluruh jawaban responden pada kuesioner. Rumus nilai akhir UAT adalah sebagai berikut.")
    add_paragraph_before(doc, before_el, "Nilai Akhir UAT = (Total Skor Aktual / Total Skor Maksimum) x 100%", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, "Keterangan:")
    for item in [
        "Total skor aktual adalah jumlah seluruh skor yang diperoleh dari semua jawaban responden pada 15 pernyataan UAT.",
        "Total skor maksimum diperoleh dari jumlah responden dikalikan jumlah pernyataan dan bobot maksimum skala Likert.",
    ]:
        add_paragraph_before(doc, before_el, item, style="List Paragraph")
    add_paragraph_before(doc, before_el, f"Total Skor Maksimum = Jumlah Responden x Jumlah Pernyataan x Bobot Maksimum", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, f"Total Skor Maksimum = {n} x 15 x 5 = {max_score}", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, "Dengan demikian, nilai akhir UAT dihitung sebagai berikut:")
    add_paragraph_before(doc, before_el, f"Nilai Akhir UAT = ({total} / {max_score}) x 100% = {final_pct}%", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_before(doc, before_el, f"\tBerdasarkan hasil perhitungan tersebut, aplikasi Sensational Glasses App memperoleh nilai akhir UAT sebesar {final_pct}% dan termasuk dalam kategori {final_cat}. Hasil ini menunjukkan bahwa aplikasi dinilai sesuai secara terbatas dari sisi pendamping siswa tunanetra dalam mendukung proses monitoring, tracking, notifikasi, peta dan navigasi, serta kontrol alarm pada perangkat Sensational Glasses.")
    add_paragraph_before(doc, before_el, "\tSelain pertanyaan tertutup berbasis skala Likert, kuesioner UAT juga dilengkapi dengan pertanyaan terbuka untuk memperoleh masukan dari responden. Ringkasan masukan terbuka dari responden ditunjukkan pada Tabel 4.14.")
    add_caption_before(doc, before_el, "Tabel 4.14 Ringkasan Masukan Terbuka Responden UAT")
    add_table_before(doc, before_el, ["No", "Aspek Pertanyaan Terbuka", "Ringkasan Jawaban Responden"], data["open_summary_rows"], widths=[0.35, 1.75, 3.85])
    add_paragraph_before(doc, before_el, "\tBerdasarkan Tabel 4.14, masukan terbuka responden menunjukkan bahwa fitur monitoring, pelacakan atau pemantauan alat, serta kontrol alarm menjadi bagian yang paling membantu. Mayoritas responden tidak menyampaikan kebutuhan fitur tambahan khusus dan menilai aplikasi sudah cukup baik. Namun, terdapat masukan bahwa bagian pengisian kredensial alat masih perlu diperjelas agar pendamping lebih mudah memahami proses menghubungkan aplikasi dengan perangkat Sensational Glasses.")
    add_paragraph_before(doc, before_el, f"\tSecara keseluruhan, hasil UAT menunjukkan bahwa Sensational Glasses App memperoleh penilaian {final_cat} dari sisi pendamping siswa tunanetra. Hasil ini memperkuat bahwa aplikasi telah sesuai secara terbatas dengan kebutuhan pendamping dalam mendukung pemantauan lokasi dan kondisi perangkat, penggunaan peta dan navigasi, penerimaan notifikasi, serta pengiriman kontrol alarm pada perangkat Sensational Glasses.")


def add_appendix(doc: Document, data):
    doc.add_page_break()
    p = doc.add_paragraph("Lampiran 9 Rekapitulasi Jawaban User Acceptance Testing (UAT)")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    doc.add_paragraph("Lampiran ini memuat rekapitulasi jawaban UAT dari guru di SLB Negeri Purwokerto. Identitas pribadi responden tidak ditampilkan dan digantikan dengan kode R1 sampai R5.")
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    table.autofit = False
    headers = ["Kode", "Peran", "Pengalaman", "Skor A1-E4", "Fitur membantu", "Bagian diperbaiki", "Fitur tambahan", "Saran lain"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in data["appendix_rows"]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=60, start=60, bottom=60, end=60)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
    widths = [0.45, 0.65, 0.85, 1.1, 1.1, 1.25, 0.9, 1.35]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    set_table_borders(table)


def main():
    data = load_uat()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC_DOCX, OUT_DOCX)
    doc = Document(OUT_DOCX)

    # Abstracts and keywords
    doc.paragraphs[55].text = (
        "Mobilitas mandiri merupakan salah satu tantangan utama bagi penyandang tunanetra, sehingga diperlukan alat bantu yang mendukung aktivitas berpindah tempat secara lebih aman dan terarah. Sensational Glasses telah dikembangkan sebagai purwarupa perangkat bantu berbasis Internet of Things (IoT), tetapi aplikasi pendukung sebelumnya masih sederhana dan belum menjadi aplikasi mobile pendamping yang terstruktur. Penelitian ini bertujuan menganalisis kebutuhan sistem aplikasi mobile pendamping, mengimplementasikan aplikasi berbasis Flutter yang terintegrasi dengan Firebase dan ThingSpeak, serta menguji aplikasi melalui Black Box Testing dan User Acceptance Testing (UAT). Metode pengembangan yang digunakan adalah prototyping melalui dua siklus. Evaluasi formatif dilakukan bersama tim PKM Sensational Glasses sebagai evaluator internal. Pengujian Black Box terhadap 56 kasus uji menghasilkan seluruh kasus valid dengan Test Case Pass Rate 100%. UAT dilakukan bersama lima guru SLB Negeri Purwokerto sebagai pendamping siswa tunanetra dan memperoleh nilai akhir 77,33% dengan kategori Baik. Hasil penelitian menunjukkan aplikasi memenuhi kebutuhan fungsional dan dinilai sesuai secara terbatas untuk mendukung monitoring, tracking, notifikasi, peta dan navigasi, serta kontrol alarm."
    )
    doc.paragraphs[56].text = "Kata Kunci: aplikasi mobile, internet of things, pelacakan GPS, Flutter, Firebase, ThingSpeak"
    doc.paragraphs[58].text = (
        "Independent mobility is one of the main challenges for people with visual impairments, requiring assistive devices that support safer and more directed movement. Sensational Glasses has been developed as an Internet of Things (IoT)-based assistive prototype, but its previous supporting application was still simple and had not become a structured companion mobile application. This study aims to analyze companion application requirements, implement a Flutter-based application integrated with Firebase and ThingSpeak, and evaluate it through Black Box Testing and User Acceptance Testing (UAT). The development method used prototyping through two cycles. Formative evaluation was conducted with the PKM Sensational Glasses team as internal evaluators. Black Box Testing on 56 test cases showed that all cases were valid, with a 100% Test Case Pass Rate. UAT was conducted with five teachers from SLB Negeri Purwokerto as companions of visually impaired students and obtained a final score of 77.33%, categorized as Good. The results indicate that the application meets functional requirements and is considered suitable in a limited context to support monitoring, tracking, notifications, maps and navigation, and alarm control."
    )
    doc.paragraphs[59].text = "Keywords: mobile application, internet of things, GPS tracking, Flutter, Firebase, ThingSpeak"

    # Replace UAT results section
    _, uat_heading = find_para(doc, "User Acceptance Testing", style="Heading 4")
    _, deploy_heading = find_para(doc, "Deploy", start=650, style="Heading 3")
    delete_between(uat_heading, deploy_heading)
    add_uat_section(doc, deploy_heading._p, data)

    # Re-read paragraph references after XML changes
    _, eval_heading = find_para(doc, "Evaluasi Kesesuaian Aplikasi Berdasarkan User Acceptance Testing", style="Heading 3")
    _, limit_heading = find_para(doc, "Keterbatasan Penelitian", style="Heading 3")
    delete_between(eval_heading, limit_heading)
    for text in [
        "User Acceptance Testing digunakan untuk menilai kesesuaian aplikasi Sensational Glasses App dari sisi pendamping siswa tunanetra setelah aplikasi dicoba secara langsung. Berbeda dengan Black Box Testing yang berfokus pada validasi fungsi aplikasi berdasarkan skenario pengujian, UAT berfokus pada penerimaan dan kesesuaian aplikasi dari sudut pandang pengguna pendamping. Pada penelitian ini, responden UAT merupakan lima guru di SLB Negeri Purwokerto yang memiliki pengalaman mengajar atau mendampingi siswa tunanetra, sehingga penilaian yang diberikan berkaitan langsung dengan kebutuhan pendampingan di lingkungan SLB Negeri Purwokerto.",
        "Berdasarkan hasil UAT pada Sub-bab 4.1.4.4, aplikasi Sensational Glasses App memperoleh nilai akhir sebesar 77,33% dan termasuk dalam kategori Baik. Nilai tersebut menunjukkan bahwa aplikasi dinilai sesuai secara terbatas oleh pendamping siswa tunanetra dalam mendukung proses monitoring dan tracking perangkat Sensational Glasses. Hasil ini juga menunjukkan bahwa aplikasi tidak hanya berjalan secara fungsional, tetapi dapat diterima dari sisi pengguna pendamping sebagai aplikasi mobile pendamping.",
        "Jika dilihat berdasarkan variabel penilaian, seluruh variabel memperoleh kategori Baik. Variabel kesesuaian dan kelayakan aplikasi pendamping memperoleh nilai tertinggi sebesar 79%, sedangkan variabel kemudahan penggunaan aplikasi memperoleh nilai 74,67%. Hal ini menunjukkan bahwa responden menilai aplikasi layak digunakan dan informasinya cukup sesuai dengan kebutuhan pendamping, tetapi aspek kemudahan penggunaan masih perlu diperhatikan agar proses awal penggunaan aplikasi menjadi lebih jelas.",
        "Variabel kejelasan informasi dan monitoring perangkat serta variabel peta, navigasi, dan notifikasi masing-masing memperoleh nilai 77,33%. Hasil ini menunjukkan bahwa informasi lokasi, status perangkat, fitur monitoring, peta, rute/navigasi, dan notifikasi dinilai membantu pendamping dalam memahami kondisi perangkat dan melakukan pelacakan. Dengan adanya informasi lokasi, kapasitas baterai, kekuatan sinyal, status perangkat, dan notifikasi kondisi penting, pendamping dapat memperoleh gambaran kondisi perangkat tanpa harus berinteraksi langsung dengan perangkat keras.",
        "Variabel kontrol alarm memperoleh nilai 78%. Hasil tersebut menunjukkan bahwa fitur kontrol alarm dinilai mudah digunakan dan dapat membantu pendamping dalam memberi respons terhadap perangkat Sensational Glasses. Masukan terbuka responden juga menunjukkan bahwa fitur pemberian perintah untuk mengaktifkan alarm menjadi salah satu fitur yang dianggap membantu.",
        "Masukan terbuka dari responden menunjukkan bahwa fitur monitoring, pelacakan atau pemantauan alat, dan kontrol alarm merupakan fitur yang paling membantu. Sebagian besar responden menyatakan tidak ada bagian yang membingungkan dan tidak mengusulkan fitur tambahan khusus. Namun, terdapat satu masukan mengenai bagian pengisian kredensial alat yang masih perlu diperjelas. Masukan tersebut menjadi catatan bahwa aplikasi perlu menyediakan penjelasan atau panduan yang lebih mudah dipahami terkait pengisian Channel ID, Read API Key, dan Write API Key.",
        "Dengan demikian, hasil UAT memperkuat hasil Black Box Testing yang telah dilakukan sebelumnya. Black Box Testing menunjukkan bahwa fungsi aplikasi berjalan valid secara teknis, sedangkan UAT menunjukkan bahwa aplikasi dapat diterima dan dinilai sesuai dari sisi pendamping siswa tunanetra dalam kategori Baik. Keterkaitan kedua hasil pengujian tersebut menunjukkan bahwa aplikasi Sensational Glasses App telah memenuhi kebutuhan fungsional dan memiliki kesesuaian penggunaan yang baik dalam konteks implementasi terbatas bersama SLB Negeri Purwokerto.",
        "Secara keseluruhan, evaluasi berdasarkan UAT menunjukkan bahwa aplikasi Sensational Glasses App telah sesuai dengan tujuan penelitian, yaitu membangun aplikasi mobile pendamping yang dapat mendukung proses monitoring, tracking, notifikasi, peta dan navigasi, serta kontrol alarm pada perangkat Sensational Glasses. Namun, karena UAT dilakukan secara terbatas kepada lima pendamping siswa tunanetra di lingkungan SLB Negeri Purwokerto, hasil ini tidak dimaksudkan untuk menggeneralisasi penerimaan seluruh pengguna secara luas, melainkan sebagai validasi terbatas terhadap kesesuaian aplikasi berdasarkan penilaian pendamping.",
    ]:
        add_paragraph_before(doc, limit_heading._p, "\t" + text)

    # Keterbatasan second paragraph: clarify five respondents
    for p in doc.paragraphs:
        if p.text.strip().startswith("Kedua, User Acceptance Testing telah dilakukan"):
            p.text = "\tKedua, User Acceptance Testing telah dilakukan secara terbatas kepada lima guru di SLB Negeri Purwokerto sebagai pendamping siswa tunanetra. Namun, jumlah responden dan durasi penggunaan masih terbatas, sehingga hasil pengujian belum dapat digeneralisasi untuk seluruh pendamping siswa tunanetra."

    # Ringkasan pembahasan paragraph containing UAT
    for p in doc.paragraphs:
        if p.text.strip().startswith("Selain pengujian fungsional, User Acceptance Testing"):
            p.text = "\tSelain pengujian fungsional, User Acceptance Testing (UAT) digunakan untuk memperoleh penilaian dari guru atau tenaga pendidik di SLB Negeri Purwokerto sebagai pendamping siswa tunanetra terhadap kesesuaian aplikasi dari sisi pendamping. Hasil UAT menunjukkan bahwa aplikasi Sensational Glasses App memperoleh nilai akhir sebesar 77,33% dan termasuk dalam kategori Baik. Dengan demikian, hasil penelitian ini menunjukkan keterhubungan antara kebutuhan sistem, implementasi aplikasi, pengujian fungsional, dan evaluasi kesesuaian aplikasi dari sisi pendamping, serta menjadi dasar penyusunan kesimpulan pada Bab V."

    # Conclusion and suggestions
    for p in doc.paragraphs:
        if "berdasarkan data simulasi UAT" in p.text:
            p.text = "Pengujian dan evaluasi aplikasi dilakukan melalui Black Box Testing dan User Acceptance Testing. Pengujian fungsional menggunakan metode Black Box Testing yang didokumentasikan dalam PDHUPL menunjukkan bahwa seluruh fitur utama aplikasi berjalan sesuai kebutuhan fungsional. Pengujian dilakukan terhadap 56 kasus uji pada modul autentikasi, monitoring perangkat, kontrol alarm, notifikasi, serta peta dan navigasi, dengan seluruh kasus uji memperoleh status valid dan Test Case Pass Rate sebesar 100%. Sementara itu, hasil UAT kepada lima guru di SLB Negeri Purwokerto sebagai pendamping siswa tunanetra menunjukkan bahwa aplikasi memperoleh nilai akhir sebesar 77,33% dan termasuk dalam kategori Baik. Hasil tersebut menunjukkan bahwa aplikasi Sensational Glasses App dinilai sesuai secara terbatas dari sisi pendamping dalam mendukung proses monitoring, tracking, notifikasi, peta dan navigasi, serta kontrol alarm."

    for p in doc.paragraphs:
        if p.text.strip() == "Menambahkan penjelasan yang lebih sederhana mengenai status perangkat seperti Aktif, Low Battery, dan Offline agar pendamping lebih mudah memahami kondisi perangkat.":
            p.text = "Menambahkan panduan yang lebih jelas mengenai pengisian kredensial perangkat, seperti Channel ID, Read API Key, dan Write API Key, agar pendamping lebih mudah menghubungkan aplikasi dengan perangkat Sensational Glasses."
        if p.text.strip() == "Menambahkan fitur pendukung seperti riwayat lokasi, panduan penggunaan awal, notifikasi suara, dan ringkasan riwayat notifikasi sesuai masukan responden UAT.":
            p.text = "Menyederhanakan istilah teknis dan alur penggunaan awal aplikasi agar pengguna pendamping yang belum terbiasa dengan konfigurasi perangkat IoT dapat memahami aplikasi dengan lebih mudah."

    add_appendix(doc, data)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
